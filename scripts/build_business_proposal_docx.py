"""Build the submission DOCX from the reviewed Markdown proposal."""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "output/business-proposal/Quotation-Desk-Business-Proposal.md"
TEMPLATE = ROOT / "hackathon_submission_materials/Quotation-Desk-Business-Proposal.docx"
BUILD_OUTPUT = ROOT / "output/business-proposal/Quotation-Desk-Business-Proposal.docx"
SUBMISSION_OUTPUT = ROOT / "hackathon_submission_materials/Quotation-Desk-Business-Proposal.docx"

NAVY = "102A43"
PALE_BLUE = "F2F6FA"
LIGHT_BORDER = "D9D9D9"
MUTED = RGBColor(76, 91, 107)


def set_font(run, name: str, size: float | None = None, color: RGBColor | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color


def add_hyperlink(paragraph, label: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5F99")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


TOKEN = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]\(https?://[^)]+\)|https?://\S+)")


def add_inline(paragraph, text: str) -> None:
    position = 0
    for match in TOKEN.finditer(text):
        if match.start() > position:
            set_font(paragraph.add_run(text[position : match.start()]), "Georgia", 10.5)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            set_font(run, "Georgia", 10.5)
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        else:
            label = token.rstrip(".,")
            add_hyperlink(paragraph, label, label)
            punctuation = token[len(label) :]
            if punctuation:
                set_font(paragraph.add_run(punctuation), "Georgia", 10.5)
        position = match.end()
    if position < len(text):
        set_font(paragraph.add_run(text[position:]), "Georgia", 10.5)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), LIGHT_BORDER)


def set_cell_padding(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table(doc: Document, rows: list[list[str]]) -> None:
    columns = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [3.35, 3.35] if columns == 2 else [1.35, 2.55, 2.8]
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        if row_index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_padding(cell)
            set_cell_borders(cell)
            shade_cell(cell, NAVY if row_index == 0 else (PALE_BLUE if row_index % 2 == 0 else "FFFFFF"))
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            add_inline(paragraph, value)
            for run in paragraph.runs:
                set_font(run, "Arial" if row_index == 0 else "Georgia", 9.1)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Georgia")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Georgia")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(25)
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(18)
    title_properties = title._element.get_or_add_pPr()
    title_borders = title_properties.find(qn("w:pBdr"))
    if title_borders is not None:
        title_properties.remove(title_borders)

    heading = doc.styles["Heading 1"]
    heading.font.name = "Arial"
    heading._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    heading._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    heading.font.size = Pt(15)
    heading.font.bold = False
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(7)
    heading.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Georgia"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Georgia")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Georgia")
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.bold = False
    caption.font.color.rgb = MUTED


def build() -> None:
    doc = Document(TEMPLATE)
    clear_body(doc)
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line == "<!-- pagebreak -->":
            doc.add_page_break()
        elif line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            add_inline(paragraph, line[2:])
            for run in paragraph.runs:
                set_font(run, "Arial", 25, RGBColor(0, 0, 0))
        elif line.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            add_inline(paragraph, line[3:])
            for run in paragraph.runs:
                set_font(run, "Arial", 15, RGBColor(0, 0, 0))
        elif line.startswith("!["):
            match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
            image_path = (SOURCE.parent / match.group(2)).resolve()
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(4)
            run = paragraph.add_run()
            picture = run.add_picture(str(image_path), width=Inches(6.9))
            picture._inline.docPr.set("descr", match.group(1))
            paragraph.paragraph_format.keep_with_next = True
        elif line.startswith("Figure "):
            paragraph = doc.add_paragraph(style="Caption")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(8)
            add_inline(paragraph, line)
            for run in paragraph.runs:
                run.bold = False
        elif line.startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            parsed = [[cell.strip() for cell in row.strip("|").split("|")] for row in table_lines]
            rows = [parsed[0], *parsed[2:]]
            add_table(doc, rows)
            continue
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline(paragraph, line[2:])
        else:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.widow_control = True
            if line.startswith("Show Me Your Agents"):
                paragraph.paragraph_format.space_after = Pt(12)
                for run in paragraph.runs:
                    set_font(run, "Georgia", 11)
            add_inline(paragraph, line)
        index += 1

    BUILD_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(BUILD_OUTPUT)
    shutil.copy2(BUILD_OUTPUT, SUBMISSION_OUTPUT)
    print(BUILD_OUTPUT)
    print(SUBMISSION_OUTPUT)


if __name__ == "__main__":
    build()
